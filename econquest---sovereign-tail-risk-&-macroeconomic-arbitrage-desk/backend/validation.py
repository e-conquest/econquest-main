"""Bounded, provider-neutral file parsing. A single bad row rejects the batch."""

import csv
import io
import json
import math
import re
import zipfile
from datetime import date, datetime, timezone
from pathlib import Path, PurePosixPath

from openpyxl import load_workbook

from .models import ANNUAL_FEATURES, COUNTRIES, Observation

MAX_UPLOAD = 5 * 1024 * 1024
MAX_ROWS = 10_000
MAX_SHEETS = 20
MAX_EXPANDED = 25 * 1024 * 1024
MAX_COLUMNS = 16
MAX_ERRORS = 50
csv.field_size_limit(65_536)


class BatchError(ValueError):
    def __init__(self, errors: list[str], row_count: int = 0, sheets: list[str] | None = None):
        super().__init__("Batch validation failed")
        self.errors = errors[:MAX_ERRORS]
        self.row_count = row_count
        self.sheets = sheets or []


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")


def timestamp(value: object, *, allow_date: bool = False) -> str:
    if isinstance(value, datetime):
        parsed = value
    else:
        text = str(value).strip()
        if allow_date and re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
            return date.fromisoformat(text).isoformat()
        if "T" not in text:
            raise ValueError("expected ISO timestamp with timezone")
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
    if parsed.tzinfo is None or parsed.utcoffset() is None:
        raise ValueError("timestamp requires a timezone")
    return parsed.astimezone(timezone.utc).isoformat().replace("+00:00", "Z")


def _missing(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _number(value: object) -> float:
    if isinstance(value, bool):
        raise ValueError("not a number")
    result = float(value)
    if not math.isfinite(result):
        raise ValueError("not finite")
    return result


def validate_annual_overlay(observations: list[Observation]) -> None:
    artifact_path = Path(__file__).resolve().parent.parent / "src" / "data" / "fittedAnnualModels.json"
    try:
        artifact = json.loads(artifact_path.read_text(encoding="utf-8"))
        countries = artifact["countries"]
        if artifact["features"] != list(ANNUAL_FEATURES):
            raise ValueError("incompatible features")
        incoming: dict[str, set[int]] = {}
        for observation in observations:
            incoming.setdefault(observation.country, set()).add(observation.year)
        for country, years in incoming.items():
            years.update(row["year"] for row in countries[country])
            ordered = sorted(years)
            if any(b != a + 1 for a, b in zip(ordered, ordered[1:])):
                raise BatchError([f"{country}: imported overlay would create a gap in bundled annual history"])
    except BatchError:
        raise
    except (OSError, KeyError, TypeError, ValueError):
        raise BatchError(["Bundled annual history is unavailable or incompatible; cannot verify overlay"]) from None


def _xlsx_rows(content: bytes, sheet: str | None) -> tuple[list[list], list[str]]:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(i.file_size for i in entries) > MAX_EXPANDED:
                raise BatchError(["XLSX expanded size or entry limit exceeded"])
            if len({entry.filename for entry in entries}) != len(entries):
                raise BatchError(["XLSX contains duplicate archive entries"])
            for entry in entries:
                path = PurePosixPath(entry.filename)
                if (
                    entry.flag_bits & 1
                    or path.is_absolute()
                    or ".." in path.parts
                    or "\\" in entry.filename
                    or entry.file_size > 10 * 1024 * 1024
                    or entry.file_size / max(entry.compress_size, 1) > 100
                ):
                    raise BatchError(["XLSX contains an unsafe or excessively compressed entry"])
        formulas = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            sheets = formulas.sheetnames
            if len(sheets) > MAX_SHEETS:
                raise BatchError(["XLSX sheet limit exceeded"])
            if not sheets:
                raise BatchError(["XLSX has no worksheets"])
            if sheet is None and len(sheets) != 1:
                raise BatchError(["Select a worksheet explicitly"], sheets=sheets)
            if sheet is not None and sheet not in sheets:
                raise BatchError(["Selected worksheet does not exist"], sheets=sheets)
            selected = sheet or sheets[0]
            ws = formulas[selected]
            if (ws.max_row or 0) > MAX_ROWS + 1 or (ws.max_column or 0) > MAX_COLUMNS:
                raise BatchError(["XLSX row or column limit exceeded"], sheets=sheets)
            # Worksheet dimensions are untrusted and may hide trailing cells/rows.
            ws.reset_dimensions()
            values = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
            try:
                values[selected].reset_dimensions()
                value_rows = values[selected].iter_rows()
                rows = []
                for row_index, formula_row in enumerate(ws.iter_rows(), 1):
                    if row_index > MAX_ROWS + 1:
                        raise BatchError(["Row limit exceeded"], MAX_ROWS + 1, sheets)
                    if len(formula_row) > MAX_COLUMNS:
                        raise BatchError(["XLSX column limit exceeded"], max(0, row_index - 1), sheets)
                    value_row = next(value_rows, ())
                    row = []
                    for col, cell in enumerate(formula_row):
                        cached = value_row[col].value if col < len(value_row) else None
                        if cell.data_type == "f" and cached is None:
                            raise BatchError(
                                [f"Row {row_index}: formula has no cached value; export values only"],
                                max(0, row_index - 1),
                                sheets,
                            )
                        row.append(cached)
                    if rows and len(row) < len(rows[0]):
                        row.extend([None] * (len(rows[0]) - len(row)))
                    rows.append(row)
                return rows, sheets
            finally:
                values.close()
        finally:
            formulas.close()
    except BatchError:
        raise
    except Exception:
        raise BatchError(["Invalid or unsupported XLSX workbook"]) from None


def parse_file(
    content: bytes, filename: str, kind: str, unit: str, sheet: str | None = None
) -> tuple[list[Observation], list[str]]:
    if len(content) > MAX_UPLOAD:
        raise BatchError(["Upload exceeds 5 MiB"])
    suffix = PurePosixPath(filename.lower()).suffix
    sheets: list[str] = []
    if suffix == ".xls":
        raise BatchError(["Legacy XLS is unsupported; export CSV or XLSX"])
    if suffix == ".xlsx":
        rows, sheets = _xlsx_rows(content, sheet)
    elif suffix == ".csv":
        if sheet:
            raise BatchError(["Worksheet selection is only valid for XLSX"])
        try:
            reader = csv.reader(io.StringIO(content.decode("utf-8-sig"), newline=""), strict=True)
            rows = []
            for row in reader:
                if len(rows) > MAX_ROWS:
                    raise BatchError(["Row limit exceeded"], MAX_ROWS + 1)
                if len(row) > MAX_COLUMNS:
                    raise BatchError(["Column limit exceeded"])
                rows.append(row)
        except (UnicodeError, csv.Error):
            raise BatchError(["CSV must be valid UTF-8 with well-formed rows"]) from None
    else:
        raise BatchError(["Only CSV and XLSX uploads are supported"])
    if not rows:
        raise BatchError(["File is empty"], sheets=sheets)
    header = [str(v).strip() if v is not None else "" for v in rows[0]]
    required = (
        ["country_code", "year", "in_default", *ANNUAL_FEATURES]
        if kind == "annual"
        else ["symbol", "observed_at", "value"]
    )
    if (
        len(set(header)) != len(header)
        or not set(required).issubset(header)
        or set(header) - set(required) - {"published_at"}
    ):
        raise BatchError(["Header must contain the documented columns without duplicates or extras"], sheets=sheets)
    if kind == "annual" and unit != "model-feature-units":
        raise BatchError(["Annual unit must be model-feature-units"], sheets=sheets)
    observations: list[Observation] = []
    errors: list[str] = []
    keys: set[tuple] = set()
    years: dict[str, list[int]] = {}
    row_count = 0
    now = datetime.now(timezone.utc)
    for line, raw in enumerate(rows[1:], 2):
        if not raw or all(_missing(v) for v in raw):
            continue
        row_count += 1
        if row_count > MAX_ROWS:
            raise BatchError(["Row limit exceeded"], row_count, sheets)
        if len(errors) >= MAX_ERRORS:
            continue
        if len(raw) != len(header):
            errors.append(f"Row {line}: column count mismatch")
            continue
        record = dict(zip(header, raw))
        try:
            published = record.get("published_at")
            published_at = None if _missing(published) else timestamp(published, allow_date=True)
            if kind == "annual":
                country = str(record["country_code"]).strip()
                if country not in COUNTRIES:
                    raise ValueError("country_code is not supported")
                year_text = str(record["year"]).strip()
                if isinstance(record["year"], bool) or not re.fullmatch(r"\d{1,4}", year_text):
                    raise ValueError("year must be an integer")
                year = int(year_text)
                if not 1900 <= year < now.year:
                    raise ValueError("year is outside the completed observation range")
                default = str(record["in_default"]).strip().lower()
                if default not in {"true", "false", "0", "1"}:
                    raise ValueError("in_default must be true, false, 0, or 1")
                features = [None if _missing(record[f]) else _number(record[f]) for f in ANNUAL_FEATURES]
                if sum(v is not None for v in features) < 4:
                    raise ValueError("at least four finite features are required")
                if any(features[i] is None for i in (2, 3, 4)):
                    raise ValueError("HMM features reservesMonths, gdpGrowth, and inflation are required")
                if any(features[i] is not None and features[i] < 0 for i in (0, 1, 2)):
                    raise ValueError("debt and reserves features must be nonnegative")
                key = (country, year)
                history = years.setdefault(country, [])
                if history and year != history[-1] + 1:
                    raise ValueError("country years must be ordered and consecutive")
                history.append(year)
                observation = Observation(
                    country=country, year=year, inDefault=default in {"true", "1"},
                    values=features, publishedAt=published_at,
                )
            else:
                symbol = str(record["symbol"]).strip()
                if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._:/^=+-]{0,63}", symbol):
                    raise ValueError("symbol must be a short identifier")
                observed = timestamp(record["observed_at"])
                if datetime.fromisoformat(observed.replace("Z", "+00:00")) > now:
                    raise ValueError("future completed observations are not allowed")
                key = (symbol, observed)
                observation = Observation(
                    symbol=symbol, observedAt=observed, value=_number(record["value"]),
                    publishedAt=published_at,
                )
            if key in keys:
                raise ValueError("duplicate observation key")
            keys.add(key)
            observations.append(observation)
        except (ValueError, TypeError, OverflowError) as exc:
            # Only fixed validation messages are exposed, never raw cell contents.
            safe = str(exc)
            allowed = (
                "country_code", "year ", "in_default", "at least", "HMM ", "debt ",
                "country years", "symbol must", "future completed", "duplicate observation",
                "timestamp requires", "expected ISO",
            )
            errors.append(f"Row {line}: {safe if safe.startswith(allowed) else 'invalid number or publication/observation date'}")
    for country, history in years.items():
        if len(history) < 2:
            errors.append(f"{country}: at least two consecutive annual observations are required")
    if not row_count:
        errors.append("At least one observation is required")
    if errors:
        raise BatchError(errors, row_count, sheets)
    if kind == "annual":
        try:
            validate_annual_overlay(observations)
        except BatchError as exc:
            raise BatchError(exc.errors, row_count, sheets) from None
    return observations, sheets
