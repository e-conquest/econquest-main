import hashlib
import io
import json
import shutil
import socket
import sqlite3
import subprocess
import sys
import time
import unittest
import urllib.error
import urllib.request
import uuid
import zipfile
from concurrent.futures import ThreadPoolExecutor
from contextlib import closing
from datetime import datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import Workbook

from backend.app import create_app
from backend.models import ANNUAL_FEATURES
from backend.storage import (
    DB_VERSION, MIGRATIONS, Conflict, Store, connect, restore_backup, verify_backup, write_snapshot,
)
from backend.validation import MAX_UPLOAD

MARKET = b"symbol,observed_at,value,published_at\nTEST,2024-01-01T00:00:00Z,100,\n"
ANNUAL_HEADER = ",".join(["country_code", "year", "in_default", *ANNUAL_FEATURES, "published_at"])
ANNUAL = (ANNUAL_HEADER + "\nCHL,2022,false,60,20,3,2,4,-1,\nCHL,2023,true,61,21,4,3,5,-2,2024-02-01\n").encode()


def workbook(rows, sheets=1) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    for row in rows:
        sheet.append(row)
    for index in range(sheets - 1):
        book.create_sheet(f"Extra{index}")
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


class BackendTests(unittest.TestCase):
    def setUp(self):
        self.work = Path("backend/.test-work") / str(uuid.uuid4())
        self.work.mkdir(parents=True)
        self.db = self.work / "research.sqlite3"
        self.app = create_app(self.db, self.work / "absent-dist")
        self.client = TestClient(self.app, base_url="http://localhost:8000", raise_server_exceptions=False)
        self.client.__enter__()
        self.token = self.client.get("/api/session").json()["token"]
        self.headers = {"X-Research-Token": self.token}

    def tearDown(self):
        self.client.__exit__(None, None, None)
        shutil.rmtree(self.work)

    def upload(self, content=MARKET, filename="data.csv", kind="market", **fields):
        fields = {"kind": kind, "source": "manual-research", "unit": "USD", **fields}
        if kind == "annual" and fields["unit"] == "USD":
            fields["unit"] = "model-feature-units"
        response = self.client.post(
            "/api/imports", headers=self.headers, data=fields,
            files={"file": (filename, content)},
        )
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def commit(self, preview):
        response = self.client.post(f"/api/imports/{preview['run']['id']}/commit", headers=self.headers)
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def test_empty_contract_and_unknown_api(self):
        self.assertEqual(self.client.get("/api/state").json(), {
            "schemaVersion": 1, "datasets": [], "runs": [], "macro": [], "weo": [], "worldbank": [], "fred": [], "bis": [], "ids": [], "comtrade": [], "gdelt": [], "acled": [], "providers": [{
                "id": "yfinance", "name": "Yahoo Finance via yfinance", "mode": "manual",
                "status": "available", "lastAttemptAt": None, "lastValidatedAt": None, "errors": [],
            }, {
                "id": "imf", "name": "Monthly IMF", "mode": "manual",
                "status": "available", "lastAttemptAt": None, "lastValidatedAt": None, "errors": [],
            }, {
                "id": "imf-weo", "name": "IMF WEO annual research and projections", "mode": "manual",
                "status": "available", "lastAttemptAt": None, "lastValidatedAt": None, "errors": [],
            }, {
                "id": "worldbank", "name": "World Bank WDI via wbgapi", "mode": "manual",
                "status": "available", "lastAttemptAt": None, "lastValidatedAt": None, "errors": [],
            }, {
                "id": "fred", "name": "FRED and ALFRED native history", "mode": "manual",
                "status": "available", "lastAttemptAt": None, "lastValidatedAt": None, "errors": [],
            }, {
                "id": "bis", "name": "BIS quarterly banking and credit", "mode": "manual",
                "status": "available", "lastAttemptAt": None, "lastValidatedAt": None, "errors": [],
            }, {
                "id": "ids", "name": "World Bank IDS via wbgapi", "mode": "manual",
                "status": "available", "lastAttemptAt": None, "lastValidatedAt": None, "errors": [],
            }, {
                "id": "comtrade", "name": "UN Comtrade via comtradeapicall", "mode": "manual",
                "status": "available", "lastAttemptAt": None, "lastValidatedAt": None, "errors": [],
            }, {
                "id": "gdelt", "name": "GDELT 2.0 daily events via gdeltPyR", "mode": "manual",
                "status": "available", "lastAttemptAt": None, "lastValidatedAt": None, "errors": [],
            }, {
                "id": "acled", "name": "ACLED core events via acled", "mode": "manual",
                "status": "available", "lastAttemptAt": None, "lastValidatedAt": None, "errors": [],
            }],
        })
        self.assertEqual(self.client.get("/api/unknown").status_code, 404)
        self.assertEqual(self.client.get("/api/session").headers["cache-control"], "no-store")

    def test_token_host_origin_and_cors(self):
        self.assertGreaterEqual(len(self.token), 32)
        for headers in ({}, {"X-Research-Token": "wrong"}, {**self.headers, "Origin": "https://evil.test"}):
            response = self.client.post("/api/imports/not-found/commit", headers=headers)
            self.assertEqual(response.status_code, 403)
        for host in ("evil.test", "localhost.evil.test", "127.0.0.1.evil.test", "localhost@evil.test", "127.1"):
            self.assertEqual(self.client.get("/api/session", headers={"Host": host}).status_code, 403)
        for origin in ("null", "http://evil.test", "http://localhost:9999"):
            self.assertEqual(self.client.get("/api/state", headers={"Origin": origin}).status_code, 403)
        for host in ("127.0.0.1:8000", "[::1]:8000"):
            self.assertEqual(self.client.get("/api/state", headers={"Host": host}).status_code, 200)
        allowed = self.client.get("/api/state", headers={"Origin": "http://localhost:5173"})
        self.assertEqual(allowed.headers["access-control-allow-origin"], "http://localhost:5173")
        self.assertEqual(self.client.get("/api/state", headers={"Origin": "http://localhost:8000"}).status_code, 200)
        preflight = self.client.options("/api/imports", headers={
            "Origin": "http://localhost:5173", "Access-Control-Request-Method": "POST",
            "Access-Control-Request-Headers": "X-Research-Token",
        })
        self.assertEqual(preflight.status_code, 200)
        self.assertEqual(self.client.options("/api/imports", headers={
            "Host": "evil.test", "Origin": "http://localhost:5173",
            "Access-Control-Request-Method": "POST",
        }).status_code, 403)

    def test_stage_commit_metadata_and_idempotency(self):
        staged = self.upload()
        self.assertEqual(staged["run"]["status"], "staged")
        self.assertIsNone(staged["run"]["datasetId"])
        self.assertIsNone(staged["preview"][0]["publishedAt"])
        self.assertEqual(self.client.get("/api/state").json()["datasets"], [])
        self.assertEqual(staged, self.upload())
        state = self.commit(staged)
        dataset = state["datasets"][0]
        self.assertEqual(dataset["version"], "1")
        self.assertTrue(dataset["active"])
        self.assertIsNone(dataset["freshnessDays"])
        self.assertEqual(len(dataset["observations"]), 1)
        self.assertNotIn("country", dataset["observations"][0])
        self.assertEqual(self.commit(staged), state)
        self.assertEqual(self.upload()["run"]["status"], "committed")
        meta = self.client.get(f"/api/datasets/{dataset['id']}/metadata").json()
        self.assertEqual(meta["originalSha256"], hashlib.sha256(MARKET).hexdigest())
        self.assertIsNone(meta["vintageAsOf"])
        self.assertFalse(meta["publicationDatesComplete"])
        self.assertTrue(meta["retrievedAt"].endswith("Z"))
        with closing(connect(self.db)) as conn:
            self.assertEqual(conn.execute("SELECT content FROM originals").fetchone()[0], MARKET)
            self.assertEqual(conn.execute("SELECT count(*) FROM runs").fetchone()[0], 1)

    def test_annual_schema_bom_missing_feature_and_publication(self):
        staged = self.upload(b"\xef\xbb\xbf" + ANNUAL.replace(b"60,20", b",20"), kind="annual")
        self.assertEqual(staged["run"]["status"], "staged")
        self.assertEqual(staged["preview"][0]["values"], [None, 20, 3, 2, 4, -1])
        self.assertIsNone(staged["preview"][0]["publishedAt"])
        self.assertEqual(staged["preview"][1]["publishedAt"], "2024-02-01")
        dataset = self.commit(staged)["datasets"][0]
        self.assertEqual(dataset["unit"], "model-feature-units")
        self.assertEqual(dataset["observations"][0]["country"], "CHL")

    def test_full_preview_and_complete_replacement_multiple_symbols(self):
        rows = b"".join(f"S{i},2024-01-01T00:00:00Z,{i}\n".encode() for i in range(150))
        preview = self.upload(b"symbol,observed_at,value\n" + rows)
        self.assertEqual(len(preview["preview"]), 150)
        state = self.commit(preview)
        self.assertEqual(len(state["datasets"][0]["observations"]), 150)
        replacement = self.commit(self.upload())
        self.assertEqual(len(replacement["datasets"]), 2)
        active = [d for d in replacement["datasets"] if d["active"]]
        self.assertEqual(len(active), 1)
        self.assertEqual(len(active[0]["observations"]), 1)

    def test_full_annual_preview_and_bundled_overlay_validation(self):
        rows = [f"CHL,{year},false,60,20,3,2,4,-1," for year in range(1900, 2025)]
        preview = self.upload((ANNUAL_HEADER + "\n" + "\n".join(rows)).encode(), kind="annual")
        self.assertEqual(preview["run"]["status"], "staged")
        self.assertEqual(len(preview["preview"]), 125)
        self.assertEqual(len(self.commit(preview)["datasets"][0]["observations"]), 125)
        gap = ANNUAL.replace(b"2022", b"1970").replace(b"2023", b"1971")
        quarantined = self.upload(gap, kind="annual")
        self.assertEqual(quarantined["run"]["status"], "quarantined")
        self.assertIn("gap", quarantined["run"]["errors"][0])

    def test_changed_bundle_quarantines_staged_annual_commit(self):
        staged = self.upload(ANNUAL, kind="annual")
        from backend.validation import BatchError
        with patch("backend.storage.validate_annual_overlay", side_effect=BatchError(["Annual overlay gap"])):
            response = self.client.post(f"/api/imports/{staged['run']['id']}/commit", headers=self.headers)
        self.assertEqual(response.status_code, 409)
        state = self.client.get("/api/state").json()
        self.assertEqual(state["runs"][0]["status"], "quarantined")
        self.assertEqual(state["datasets"], [])

    def test_all_countries_and_annual_validation(self):
        invalid = [
            ANNUAL.replace(b"CHL", b"USA"),
            ANNUAL.replace(b"2023", b"2024"),
            ANNUAL.replace(b"2022", b"2023"),
            ANNUAL.replace(b"false", b"maybe"),
            ANNUAL.replace(b"60,20,3,2,4,-1", b"-1,20,3,2,4,-1"),
            ANNUAL.replace(b"60,20,3,2,4,-1", b"60,20,,2,4,-1"),
            ANNUAL.replace(b"60,20,3,2,4,-1", b",,3,2,4,"),
            ANNUAL.replace(b"60,20", b"NaN,20"),
            ANNUAL.replace(b"60,20", b"inf,20"),
            ANNUAL.splitlines()[0] + b"\n" + ANNUAL.splitlines()[1] + b"\n",
            ANNUAL.replace(b"2023", str(datetime.now(timezone.utc).year).encode()),
        ]
        for content in invalid:
            with self.subTest(content=content):
                preview = self.upload(content, kind="annual")
                self.assertEqual(preview["run"]["status"], "quarantined")
                self.assertEqual(preview["preview"], [])
                self.assertTrue(preview["run"]["errors"])
                self.assertEqual(
                    self.client.post(f"/api/imports/{preview['run']['id']}/commit", headers=self.headers).status_code, 409
                )
        self.assertEqual(self.client.get("/api/state").json()["datasets"], [])
        for country in ("CHL", "BLZ", "LKA", "ARG", "TUR", "GHA", "IDN", "PAK", "ZAF"):
            self.assertEqual(self.upload(ANNUAL.replace(b"CHL", country.encode()), kind="annual")["run"]["status"], "staged")

    def test_annual_order_and_unit(self):
        lines = ANNUAL.splitlines()
        self.assertEqual(self.upload(b"\n".join([lines[0], lines[2], lines[1]]), kind="annual")["run"]["status"], "quarantined")
        self.assertEqual(self.upload(ANNUAL, kind="annual", unit="percent")["run"]["status"], "quarantined")

    def test_market_rejects_entire_invalid_batches(self):
        invalid = [
            b"symbol,observed_at,value\n",
            MARKET + MARKET.splitlines()[1] + b"\n",
            MARKET.replace(b",100,", b",NaN,"),
            MARKET.replace(b",100,", b",inf,"),
            MARKET.replace(b",100,", b",,"),
            MARKET.replace(b"2024-01-01T00:00:00Z", b"2024-01-01T00:00:00"),
            MARKET.replace(b"2024-01-01T00:00:00Z", b"9999-01-01T00:00:00Z"),
            MARKET.replace(b"TEST", b"bad symbol"),
            MARKET.replace(b",100,", b",100,2020-99-99"),
            b"symbol,observed_at,value,value\nTEST,2020-01-01T00:00:00Z,1,2\n",
            b"symbol,observed_at,value\nTEST,2020-01-01T00:00:00Z,1,unexpected\n",
            b"\xff\xfeinvalid",
            b'symbol,observed_at,value\n"unterminated',
        ]
        for content in invalid:
            with self.subTest(content=content):
                result = self.upload(content)
                self.assertEqual(result["run"]["status"], "quarantined")
                self.assertEqual(result["preview"], [])
        self.assertEqual(self.client.get("/api/state").json()["datasets"], [])

    def test_timezone_normalization_duplicate_key(self):
        content = b"symbol,observed_at,value\nTEST,2024-01-01T01:00:00+01:00,1\n"
        self.assertEqual(self.upload(content)["preview"][0]["observedAt"], "2024-01-01T00:00:00Z")
        duplicate = content + b"TEST,2024-01-01T00:00:00Z,2\n"
        self.assertEqual(self.upload(duplicate)["run"]["status"], "quarantined")

    def test_revision_activation_and_rollback_history(self):
        first = self.commit(self.upload())["datasets"][0]
        second = self.commit(self.upload(MARKET.replace(b",100,", b",101,"), source="second-source"))["datasets"][1]
        self.assertEqual(second["version"], "2")
        state = self.client.post(f"/api/datasets/{first['id']}/activate", headers=self.headers).json()
        self.assertTrue(state["datasets"][0]["active"])
        self.assertFalse(state["datasets"][1]["active"])
        self.client.post(f"/api/datasets/{second['id']}/activate", headers=self.headers)
        rolled = self.client.post(f"/api/datasets/{second['id']}/rollback", headers=self.headers)
        self.assertEqual(rolled.status_code, 200)
        state = rolled.json()
        self.assertEqual(state["runs"][1]["status"], "rolled_back")
        self.assertTrue(state["datasets"][0]["active"])
        self.assertFalse(state["datasets"][1]["active"])
        self.assertEqual(len(state["datasets"][1]["observations"]), 1)
        self.assertEqual(self.client.post(f"/api/datasets/{second['id']}/activate", headers=self.headers).status_code, 409)
        events = self.client.get("/api/audit").json()["events"]
        self.assertIn("rolled_back", [e["event"] for e in events])
        count = len(events)
        self.client.post(f"/api/datasets/{second['id']}/rollback", headers=self.headers)
        self.assertEqual(len(self.client.get("/api/audit").json()["events"]), count)
        third = self.commit(self.upload(MARKET.replace(b",100,", b",102,")))["datasets"][2]
        self.assertEqual(third["version"], "3")

    def test_one_active_dataset_each_kind_and_no_source_mixing(self):
        self.commit(self.upload())
        self.commit(self.upload(ANNUAL, kind="annual"))
        self.commit(self.upload(MARKET.replace(b"TEST", b"OTHER"), source="other"))
        active = [d for d in self.client.get("/api/state").json()["datasets"] if d["active"]]
        self.assertEqual({d["kind"] for d in active}, {"annual", "market"})
        market = next(d for d in active if d["kind"] == "market")
        self.assertEqual([o["symbol"] for o in market["observations"]], ["OTHER"])

    def test_staged_quarantine_committed_and_health_survive_restart(self):
        staged = self.upload()
        quarantine = self.upload(b"bad")
        self.commit(self.upload(ANNUAL, kind="annual"))
        before = self.client.get("/api/state").json()
        self.client.get("/api/health")
        self.client.__exit__(None, None, None)
        self.app = create_app(self.db)
        self.client = TestClient(self.app, base_url="http://localhost:8000", raise_server_exceptions=False)
        self.client.__enter__()
        self.assertEqual(self.client.get("/api/state").json(), before)
        self.assertEqual(self.client.get("/api/session").json()["token"], self.token)
        self.assertEqual(next(r for r in before["runs"] if r["id"] == staged["run"]["id"])["status"], "staged")
        self.assertEqual(next(r for r in before["runs"] if r["id"] == quarantine["run"]["id"])["status"], "quarantined")
        self.assertIn("dataset:", json.dumps(self.client.get("/api/health").json()))
        self.commit(staged)

    def test_health_unknown_freshness_and_transition_event_dedup(self):
        self.commit(self.upload())
        health = self.client.get("/api/health").json()
        data_health = next(e for e in health["observations"] if e["id"].startswith("dataset:"))
        self.assertEqual(data_health["status"], "unknown")
        self.assertIsNone(data_health["details"]["freshnessDays"])
        self.commit(self.upload(freshnessDays="1"))
        health = self.client.get("/api/health").json()
        self.assertIn("stale", [e["status"] for e in health["observations"]])
        count = len(self.client.get("/api/events").json()["events"])
        self.client.get("/api/health")
        self.client.get("/api/health")
        self.assertEqual(len(self.client.get("/api/events").json()["events"]), count)
        events = self.client.get("/api/events").json()["events"]
        transitions = [e for e in events if e["event"] == "data_health_changed"]
        self.assertTrue(transitions)
        self.assertTrue(all(e["createdAt"].endswith("Z") for e in transitions))
        self.assertFalse(health["connectorsActive"])

    def test_health_uses_oldest_latest_symbol_not_globally_newest(self):
        now = datetime.now(timezone.utc)
        recent = (now - timedelta(hours=1)).isoformat()
        old = (now - timedelta(days=10)).isoformat()
        content = (
            "symbol,observed_at,value\n"
            f"FRESH,{old},1\nFRESH,{recent},2\nSTALE,{old},3\n"
        ).encode()
        dataset = self.commit(self.upload(content, freshnessDays="2"))["datasets"][0]
        health = next(
            entry for entry in self.client.get("/api/health").json()["observations"]
            if entry["id"] == f"dataset:{dataset['id']}"
        )
        self.assertEqual(health["status"], "stale")
        self.assertGreater(health["details"]["ageDays"], 9)
        self.assertEqual(len(health["details"]["seriesLatestObservationAt"]), 2)
        events = self.client.get("/api/events").json()["events"]
        self.assertTrue(any(
            event["event"] == "data_health_changed"
            and event["details"].get("healthId") == health["id"]
            and event["details"]["status"] == "stale" for event in events
        ))
        count = len(events)
        self.client.get("/api/health")
        self.assertEqual(len(self.client.get("/api/events").json()["events"]), count)

    def test_health_uses_oldest_latest_country_completed_year_not_publication(self):
        year = datetime.now(timezone.utc).year
        rows = [
            f"CHL,{year - 2},false,60,20,3,2,4,-1,",
            f"CHL,{year - 1},false,60,20,3,2,4,-1,",
            # A recent publication must not conceal an old observation period.
            f"BLZ,{year - 4},false,60,20,3,2,4,-1,{year}-01-01",
            f"BLZ,{year - 3},false,60,20,3,2,4,-1,{year}-01-01",
        ]
        content = (ANNUAL_HEADER + "\n" + "\n".join(rows)).encode()
        dataset = self.commit(self.upload(content, kind="annual", freshnessDays="400"))["datasets"][0]
        health = next(
            entry for entry in self.client.get("/api/health").json()["observations"]
            if entry["id"] == f"dataset:{dataset['id']}"
        )
        self.assertEqual(health["status"], "stale")
        self.assertGreater(health["details"]["ageDays"], 700)
        dates = health["details"]["seriesLatestObservationAt"]
        self.assertEqual(dates["CHL"], f"{year - 1}-12-31T23:59:59Z")
        self.assertEqual(dates["BLZ"], f"{year - 3}-12-31T23:59:59Z")

    def test_failed_commit_is_atomic_and_records_interrupted_run(self):
        staged = self.upload()
        with patch("backend.storage._audit", side_effect=RuntimeError("secret-cell-value")):
            # Simulate a failure after dataset and observation inserts, before commit.
            with self.assertRaises(RuntimeError):
                self.app.state.store.commit(staged["run"]["id"])
        self.assertEqual(self.client.get("/api/state").json()["datasets"], [])
        self.app.state.store.record_failure(staged["run"]["id"])
        state = self.client.get("/api/state").json()
        self.assertEqual(state["runs"][0]["status"], "interrupted")
        self.assertNotIn("secret-cell-value", json.dumps(state))
        self.assertEqual(self.client.post(f"/api/imports/{staged['run']['id']}/commit", headers=self.headers).status_code, 409)

    def test_api_failed_commit_records_error(self):
        staged = self.upload()
        with self.assertLogs("research.backend", level="INFO") as logs, patch.object(
            self.app.state.store, "commit", side_effect=sqlite3.OperationalError("sensitive-path")
        ):
            response = self.client.post(f"/api/imports/{staged['run']['id']}/commit", headers=self.headers)
        self.assertEqual(response.status_code, 500)
        self.assertNotIn("sensitive-path", response.text)
        self.assertNotIn("sensitive-path", "\n".join(logs.output))
        self.assertIn("OperationalError", "\n".join(logs.output))
        self.assertEqual(self.client.get("/api/state").json()["runs"][0]["status"], "interrupted")

    def test_metadata_validation_and_unknown_resources(self):
        for fields in ({"kind": "other"}, {"source": ""}, {"unit": ""}, {"freshnessDays": "-1"}, {"freshnessDays": "1.5"}):
            response = self.client.post(
                "/api/imports", headers=self.headers,
                data={"kind": "market", "source": "s", "unit": "u", **fields},
                files={"file": ("x.csv", MARKET)},
            )
            self.assertEqual(response.status_code, 422, response.text)
        for path in ("/api/imports/unknown/commit", "/api/datasets/unknown/activate", "/api/datasets/unknown/rollback"):
            self.assertEqual(self.client.post(path, headers=self.headers).status_code, 404)
        self.assertEqual(self.client.get("/api/datasets/unknown/metadata").status_code, 404)

    def test_csv_row_upload_and_field_limits(self):
        content = b"symbol,observed_at,value\n" + b"T,2020-01-01T00:00:00Z,1\n" * 10001
        result = self.upload(content)
        self.assertEqual(result["run"]["status"], "quarantined")
        self.assertIn("limit", result["run"]["errors"][0].lower())
        result = self.upload(b"x" * (MAX_UPLOAD + 1))
        self.assertEqual(result["run"]["status"], "quarantined")
        self.assertIn("5 MiB", result["run"]["errors"][0])
        response = self.client.post(
            "/api/imports", headers=self.headers, files={"file": ("large.csv", b"x" * (MAX_UPLOAD + 100000))},
            data={"kind": "market", "source": "s", "unit": "u"},
        )
        self.assertEqual(response.status_code, 413)
        self.assertEqual(self.upload(b"symbol,observed_at,value\n" + b"a" * 70000)["run"]["status"], "quarantined")

    def test_xlsx_values_sheets_and_formula_rejection(self):
        rows = [["symbol", "observed_at", "value"], ["X", "2024-01-01T00:00:00Z", 5]]
        result = self.upload(workbook(rows), "book.xlsx")
        self.assertEqual(result["sheets"], ["Data"])
        self.assertEqual(result["preview"][0]["value"], 5)
        multiple = workbook(rows, sheets=2)
        result = self.upload(multiple, "book.xlsx")
        self.assertEqual(result["run"]["status"], "quarantined")
        self.assertEqual(result["sheets"], ["Data", "Extra0"])
        self.assertEqual(self.upload(multiple, "book.xlsx", sheet="Data")["run"]["status"], "staged")
        self.assertEqual(self.upload(multiple, "book.xlsx", sheet="Missing")["run"]["status"], "quarantined")
        formula_rows = [rows[0], ["X", "2024-01-01T00:00:00Z", "=2+3"]]
        formula = self.upload(workbook(formula_rows), "formulas.xlsx")
        self.assertEqual(formula["run"]["status"], "quarantined")
        self.assertIn("cached value", formula["run"]["errors"][0])
        self.assertEqual(self.upload(b"unsupported", "data.xls")["run"]["status"], "quarantined")
        self.assertEqual(self.upload(b"not-a-zip", "data.xlsx")["run"]["status"], "quarantined")
        self.assertEqual(self.upload(workbook(rows, sheets=21), "book.xlsx")["run"]["status"], "quarantined")

    def test_xlsx_zip_expansion_and_dimensions(self):
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("huge.xml", b"a" * (1024 * 1024))
        result = self.upload(output.getvalue(), "bomb.xlsx")
        self.assertEqual(result["run"]["status"], "quarantined")
        self.assertIn("compressed", result["run"]["errors"][0])
        book = Workbook()
        book.active.cell(row=10002, column=1, value="bad")
        output = io.BytesIO()
        book.save(output)
        book.close()
        result = self.upload(output.getvalue(), "dimensions.xlsx")
        self.assertEqual(result["run"]["status"], "quarantined")
        self.assertIn("limit", result["run"]["errors"][0])

    def test_xlsx_misleading_dimensions_do_not_hide_invalid_rows(self):
        content = workbook([
            ["symbol", "observed_at", "value", "published_at"],
            ["X", "2024-01-01T00:00:00Z", 5, None],
        ])
        self.assertEqual(self.upload(content, "blank-publication.xlsx")["run"]["status"], "staged")
        source = workbook([
            ["symbol", "observed_at", "value"],
            ["X", "2024-01-01T00:00:00Z", 5],
            ["X", "2024-01-01T00:00:00Z", 6],
        ])
        output = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(source)) as archive, zipfile.ZipFile(output, "w") as target:
            for entry in archive.infolist():
                data = archive.read(entry.filename)
                if entry.filename == "xl/worksheets/sheet1.xml":
                    data = data.replace(b'<dimension ref="A1:C3"/>', b'<dimension ref="A1:C2"/>')
                target.writestr(entry, data)
        result = self.upload(output.getvalue(), "misleading.xlsx")
        self.assertEqual(result["run"]["status"], "quarantined")
        self.assertIn("duplicate", result["run"]["errors"][0])

    def test_originals_observations_and_audit_are_immutable(self):
        self.commit(self.upload())
        with closing(connect(self.db)) as conn:
            for table in ("originals", "observations", "audit"):
                with self.assertRaises(sqlite3.IntegrityError):
                    conn.execute(f"DELETE FROM {table}")
            with self.assertRaises(sqlite3.IntegrityError):
                conn.execute("UPDATE originals SET content=x'00'")
            with self.assertRaises(sqlite3.IntegrityError):
                conn.execute("UPDATE observations SET payload='{}'")

    def test_backup_restore_integrity_and_no_overwrite(self):
        self.commit(self.upload())
        self.upload(ANNUAL, kind="annual")
        response = self.client.get("/api/backup")
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment", response.headers["content-disposition"])
        snapshot = self.work / "snapshot.sqlite3"
        snapshot.write_bytes(response.content)
        verify_backup(snapshot)
        restored = self.work / "restored.sqlite3"
        with self.assertRaises(Conflict):
            restore_backup(snapshot, restored, offline=False)
        restore_backup(snapshot, restored, offline=True)
        with self.assertRaises(Conflict):
            restore_backup(snapshot, restored, offline=True)
        with self.assertRaises(Conflict):
            restore_backup(snapshot, self.db, offline=True)
        restored_store = Store(restored)
        restored_store.open()
        try:
            self.assertEqual(restored_store.state(), self.app.state.store.state())
        finally:
            restored_store.close()
        second = self.work / "cli-backup.sqlite3"
        result = subprocess.run(
            [sys.executable, "-m", "backend", "backup", "--db", str(self.db), "--output", str(second)],
            capture_output=True, text=True,
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        verify_backup(second)
        result = subprocess.run(
            [sys.executable, "-m", "backend", "restore", "--input", str(second),
             "--target", str(self.work / "cli-restored.sqlite3"), "--offline"],
            capture_output=True, text=True,
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        corrupt = self.work / "corrupt.sqlite3"
        corrupt.write_bytes(b"not a database")
        with self.assertRaises(sqlite3.DatabaseError):
            restore_backup(corrupt, self.work / "bad-target.sqlite3", offline=True)
        self.assertFalse((self.work / "bad-target.sqlite3").exists())

    def test_migration_from_v1_and_reject_future_schema(self):
        legacy = self.work / "legacy.sqlite3"
        with closing(connect(legacy)) as conn:
            for statement in MIGRATIONS[1]:
                conn.execute(statement)
            conn.execute("PRAGMA user_version=1")
        store = Store(legacy)
        store.open()
        try:
            self.assertEqual(store.state().datasets, [])
            self.assertEqual(store.health()["status"], "ok")
            with closing(connect(legacy)) as conn:
                self.assertEqual(conn.execute("PRAGMA user_version").fetchone()[0], DB_VERSION)
        finally:
            store.close()
        with closing(connect(legacy)) as conn:
            conn.execute("PRAGMA user_version=999")
        with self.assertRaises(Conflict):
            store.open()
        self.assertIsNone(store.lock.handle)

    def test_concurrent_idempotent_staging_commit_and_process_lock(self):
        with ThreadPoolExecutor(max_workers=4) as executor:
            results = list(executor.map(lambda _: self.upload(), range(4)))
        self.assertEqual(len({r["run"]["id"] for r in results}), 1)
        with ThreadPoolExecutor(max_workers=4) as executor:
            states = list(executor.map(lambda _: self.commit(results[0]), range(4)))
        self.assertTrue(all(len(s["datasets"]) == 1 for s in states))
        other = Store(self.db)
        with self.assertRaises(Conflict):
            other.open()

    def test_migration_preserves_originals_revisions_and_durable_jobs(self):
        self.commit(self.upload())
        self.upload(ANNUAL, kind="annual")
        self.upload(b"invalid")
        before = self.app.state.store.state()
        legacy = self.work / "populated-v1.sqlite3"
        legacy.write_bytes(self.app.state.store.snapshot())
        with closing(connect(legacy)) as conn:
            conn.execute("DROP TABLE health")
            for table in ("analysis_bundles", "acled_rows", "acled_chunks", "acled_sources", "acled_runs",
                          "gdelt_rows", "gdelt_chunks", "gdelt_sources", "gdelt_runs",
                          "comtrade_rows", "comtrade_sources", "comtrade_tasks", "comtrade_requests", "comtrade_runs"):
                conn.execute(f"DROP TABLE {table}")
            for table in ("ids_rows", "ids_series", "ids_sources", "ids_runs", "bis_rows", "bis_series", "bis_sources", "bis_runs"):
                conn.execute(f"DROP TABLE {table}")
            conn.execute("DROP TABLE fred_rows")
            conn.execute("DROP TABLE fred_series")
            conn.execute("DROP TABLE fred_sources")
            conn.execute("DROP TABLE fred_runs")
            conn.execute("DROP TABLE worldbank_rows")
            conn.execute("DROP TABLE worldbank_series")
            conn.execute("DROP TABLE worldbank_sources")
            conn.execute("DROP TABLE worldbank_runs")
            conn.execute("DROP TABLE weo_rows")
            conn.execute("DROP TABLE weo_series")
            conn.execute("DROP TABLE weo_sources")
            conn.execute("DROP TABLE weo_runs")
            conn.execute("DROP TABLE macro_coverage")
            conn.execute("DROP TABLE macro_rows")
            conn.execute("DROP TABLE macro_sources")
            conn.execute("DROP TABLE macro_runs")
            conn.execute("DROP TABLE provider_pulls")
            conn.execute("PRAGMA user_version=1")
        store = Store(legacy)
        store.open()
        try:
            self.assertEqual(store.state(), before)
            with closing(connect(legacy)) as conn:
                self.assertEqual(conn.execute("SELECT count(*) FROM originals").fetchone()[0], 3)
            verify_backup(legacy)
        finally:
            store.close()

    def test_static_spa_never_catches_api_or_missing_assets(self):
        dist = self.work / "dist"
        dist.mkdir()
        (dist / "index.html").write_text("<html>Research dashboard</html>")
        (dist / "asset.js").write_text("export const test = true;")
        with TestClient(create_app(self.work / "static.sqlite3", dist), base_url="http://localhost") as client:
            self.assertIn("Research dashboard", client.get("/research").text)
            self.assertEqual(client.get("/asset.js").status_code, 200)
            self.assertEqual(client.get("/missing.js").status_code, 404)
            self.assertEqual(client.get("/api/missing").status_code, 404)
            self.assertEqual(client.get("/api").status_code, 404)

    def test_environment_db_path_and_loopback_origin_configuration(self):
        configured_db = self.work / "configured.sqlite3"
        with patch.dict("os.environ", {
            "ECONQUEST_DB_PATH": str(configured_db),
            "ECONQUEST_ALLOWED_ORIGINS": "http://127.0.0.1:4179,http://localhost:8123",
        }):
            with TestClient(create_app(), base_url="http://localhost") as client:
                self.assertTrue(configured_db.exists())
                allowed = client.get("/api/state", headers={"Origin": "http://127.0.0.1:4179"})
                self.assertEqual(allowed.status_code, 200)
                self.assertEqual(allowed.headers["access-control-allow-origin"], "http://127.0.0.1:4179")
                self.assertEqual(client.get("/api/state", headers={"Origin": "http://localhost:5173"}).status_code, 403)
        for invalid in ("*", "https://evil.test", "http://localhost:1234/path", "null"):
            with patch.dict("os.environ", {"ECONQUEST_ALLOWED_ORIGINS": invalid}):
                with self.assertRaises(ValueError):
                    create_app(self.work / "not-created.sqlite3")

    def test_actual_cli_http_runtime(self):
        with socket.socket() as sock:
            sock.bind(("127.0.0.1", 0))
            port = sock.getsockname()[1]
        process = subprocess.Popen(
            [sys.executable, "-m", "backend", "serve", "--db", str(self.work / "runtime.sqlite3"),
             "--port", str(port)],
            stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True,
        )
        try:
            deadline = time.monotonic() + 15
            while True:
                try:
                    with urllib.request.urlopen(f"http://127.0.0.1:{port}/api/state", timeout=1) as response:
                        payload = json.load(response)
                    break
                except (urllib.error.URLError, TimeoutError):
                    if time.monotonic() > deadline or process.poll() is not None:
                        self.fail("CLI server failed to become responsive")
                    time.sleep(0.1)
            self.assertEqual(payload["schemaVersion"], 1)
            self.assertEqual(payload["providers"][0]["id"], "yfinance")
            self.assertEqual(payload["providers"][0]["status"], "available")
        finally:
            process.terminate()
            process.communicate(timeout=10)


if __name__ == "__main__":
    unittest.main()
